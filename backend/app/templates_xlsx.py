"""Downloadable roster and marks templates (spec §6, §8 item 2).

Generated from the teacher's current question list rather than a fixed default, so the
marks template always has exactly the columns the batch expects. This is the backend
copy referred to in the spec's §6; the frontend downloads from here rather than building
its own with SheetJS, so the two cannot drift apart.
"""

from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import ClassInfo

_HEADER_FONT = Font(bold=True)
_HEADER_FILL = PatternFill("solid", fgColor="DDDDDD")
_CENTRE = Alignment(horizontal="center", vertical="center")


def _write_header(sheet, headers: list[str], widths: list[int]) -> None:
    for index, (title, width) in enumerate(zip(headers, widths), start=1):
        cell = sheet.cell(row=1, column=index, value=title)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _CENTRE
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"


def build_roster_template() -> bytes:
    """Roll + Name, with two example rows (spec §4.2)."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Roster"

    _write_header(sheet, ["Roll", "Name"], [18, 32])
    sheet["A1"].comment = Comment(
        "Roll number as plain text. This is the only key linking the roster, the "
        "marks sheet and the signature form - it must match exactly.",
        "Top Sheet Generator",
    )
    for row, (roll, name) in enumerate(
        [("25371025001", "Abhisha Banerjee"), ("25371025003", "Angira Ghosh")], start=2
    ):
        sheet.cell(row=row, column=1, value=roll)
        sheet.cell(row=row, column=2, value=name)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def build_marks_template(class_info: ClassInfo) -> bytes:
    """Roll + Name + one column per configured question (spec §4.3).

    Question numbers are written as **text**, not numbers. Left as numbers, Excel stores
    question "2" as the float 2.0 -- which is exactly the header shape that broke score
    matching in the real working file (Appendix A.9).
    """
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Marks"

    headers = ["Roll", "Name"] + [q.qno for q in class_info.questions]
    widths = [18, 28] + [8] * len(class_info.questions)
    _write_header(sheet, headers, widths)

    for index, question in enumerate(class_info.questions, start=3):
        cell = sheet.cell(row=1, column=index)
        cell.value = str(question.qno)
        cell.number_format = "@"  # text, so "2" cannot become 2.0
        cell.comment = Comment(
            f"Question {question.qno} - out of {question.marks_allotted:g} marks.\n"
            f"Enter the marks awarded, or leave blank if not attempted.\n"
            f"Blank is not the same as 0: blank prints as 'NA', 0 prints as 'Wrong'.",
            "Top Sheet Generator",
        )

    example = ["25371025001", "Abhisha Banerjee"] + [
        1 if q.marks_allotted == 1 else 3 for q in class_info.questions
    ]
    for column, value in enumerate(example, start=1):
        sheet.cell(row=2, column=column, value=value)

    notes = workbook.create_sheet("Notes")
    notes.column_dimensions["A"].width = 100
    lines = [
        "How to fill in this sheet",
        "",
        "1. One row per student. 'Roll' must match the roster and the signature form exactly.",
        "2. Enter the marks awarded for each question. Half marks (0.5) are allowed.",
        "3. Leave a cell BLANK if the question was not attempted or does not apply.",
        "   Blank prints as 'NA'. A literal 0 prints as 'Wrong'. They are not the same.",
        "4. Do not add a Total column - the total is computed, and any typed total is ignored.",
        "",
        "How the total is calculated for this paper:",
    ]
    for group in class_info.question_groups:
        lines.append(
            f"   - {group.label or 'Group'}: best {group.count_best} of "
            f"{len(group.questions)} ({', '.join(group.questions)})"
        )
    if not class_info.question_groups:
        lines.append("   - Every question counts in full.")
    lines += ["", f"Full marks: {class_info.full_marks:g}"]

    for row, line in enumerate(lines, start=1):
        cell = notes.cell(row=row, column=1, value=line)
        if row == 1:
            cell.font = Font(bold=True, size=13)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
